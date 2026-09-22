import streamlit as st
import datetime
import calendar
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

st.set_page_config(page_title="SUSUZ HES SCADA Paneli", layout="wide", page_icon="⚡")
st.title("⚡ SUSUZ HES İşletme & SCADA Kontrol Paneli")

saatler_listesi = [f"{str(i).zfill(2)}:00" for i in range(24)]
if "saat_idx" not in st.session_state:
    st.session_state["saat_idx"] = 0

AYLAR_TR = {
    1: "OCAK", 2: "SUBAT", 3: "MART", 4: "NISAN", 5: "MAYIS", 6: "HAZIRAN",
    7: "TEMMUZ", 8: "AGUSTOS", 9: "EYLUL", 10: "EKIM", 11: "KASIM", 12: "ARALIK"
}

# --- 1. VARDİYA VE ZAMAN SEÇİMİ ---
c_t1, c_t2, c_t3 = st.columns([1, 1, 2])
with c_t1:
    tarih = st.date_input("📅 Tarih", datetime.date.today())
with c_t2:
    secilen_saat = st.selectbox("⏰ Saat", saatler_listesi, index=st.session_state["saat_idx"])
with c_t3:
    sorumlu = st.text_input("👷 Vardiya Sorumlusu", placeholder="Örn: Murat / Çağlın")

# Dosya adı: O ayın tek çalışma kitabı (Örn: EYLUL_2026.xlsx)
ay_adi = AYLAR_TR[tarih.month]
EXCEL_FILE = f"{ay_adi}_{tarih.year}.xlsx"
sayfa_adi = f"{str(tarih.day).zfill(2)}"  # Sekme adı: 01, 02, ..., 22, 23 gibi

def bos_gunluk_sablon_ciz(ws):
    ws.views.sheetView[0].showGridLines = True
    
    font_ana = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_alt = Font(name="Arial", size=8, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_ince = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    f_var = PatternFill("solid", fgColor="34495E")
    f_u1 = PatternFill("solid", fgColor="1F4E79")
    f_u2 = PatternFill("solid", fgColor="2CA02C")
    f_u3 = PatternFill("solid", fgColor="C00000")
    f_reg = PatternFill("solid", fgColor="7030A0")
    
    # 1. Satır: Bloklar
    ws.merge_cells("A1:B1"); ws["A1"] = "VARDİYA"; ws["A1"].fill = f_var; ws["A1"].font = font_ana; ws["A1"].alignment = align_center
    ws.merge_cells("C1:J1"); ws["C1"] = "ÜNİTE 1 (3100 kW)"; ws["C1"].fill = f_u1; ws["C1"].font = font_ana; ws["C1"].alignment = align_center
    ws.merge_cells("K1:R1"); ws["K1"] = "ÜNİTE 2 (3100 kW)"; ws["K1"].fill = f_u2; ws["K1"].font = font_ana; ws["K1"].alignment = align_center
    ws.merge_cells("S1:Z1"); ws["S1"] = "ÜNİTE 3 (1200 kW)"; ws["S1"].fill = f_u3; ws["S1"].font = font_ana; ws["S1"].alignment = align_center
    ws.merge_cells("AA1:AD1"); ws["AA1"] = "REGÜLATÖR & SAYAÇ"; ws["AA1"].fill = f_reg; ws["AA1"].font = font_ana; ws["AA1"].alignment = align_center

    # 2. Satır: Alt Başlıklar
    headers = [
        ("Vardiyacı", f_var), ("Saat", f_var),
        ("U1 Aktif", f_u1), ("U1 Reaktif", f_u1), ("U1 Volt", f_u1), ("U1 Akım", f_u1), ("U1 Kanat", f_u1), ("U1 Ön-1", f_u1), ("U1 Ön-2", f_u1), ("U1 Arka", f_u1),
        ("U2 Aktif", f_u2), ("U2 Reaktif", f_u2), ("U2 Volt", f_u2), ("U2 Akım", f_u2), ("U2 Kanat", f_u2), ("U2 Ön-1", f_u2), ("U2 Ön-2", f_u2), ("U2 Arka", f_u2),
        ("U3 Aktif", f_u3), ("U3 Reaktif", f_u3), ("U3 Volt", f_u3), ("U3 Akım", f_u3), ("U3 Kanat", f_u3), ("U3 Ön-1", f_u3), ("U3 Ön-2", f_u3), ("U3 Arka", f_u3),
        ("Havuz Kot", f_reg), ("Havuz Fark", f_reg), ("Anasayaç", f_reg), ("Saatlik Fark", f_reg)
    ]
    for c_idx, (txt, fill_renk) in enumerate(headers, start=1):
        c = ws.cell(row=2, column=c_idx, value=txt)
        c.fill = fill_renk; c.font = font_alt; c.alignment = align_center; c.border = border_ince

    # 24 Saati Satır Satır Diz
    for h in range(24):
        r = 3 + h
        s_cell = ws.cell(row=r, column=2, value=f"{str(h).zfill(2)}:00")
        s_cell.alignment = align_center
        s_cell.font = Font(name="Arial", size=9, bold=True)
        for c in range(1, 31):
            ws.cell(row=r, column=c).border = border_ince

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 11

def calisma_kitabini_hazirla(yil, ay):
    wb = openpyxl.Workbook()
    # Varsayılan boş sayfayı sil
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    
    # O ay kaç gün çekiyorsa (örneğin Eylül için 30 gün) hepsinin sayfasını aç
    toplam_gun = calendar.monthrange(yil, ay)[1]
    for g in range(1, toplam_gun + 1):
        s_name = str(g).zfill(2)
        ws = wb.create_sheet(title=s_name)
        bos_gunluk_sablon_ciz(ws)
    return wb

def veriyi_excele_yaz(secilen_tarih, saat_str, sorumlu, u1_vals, u2_vals, u3_vals, havuz, sayac):
    yil = secilen_tarih.year
    ay = secilen_tarih.month
    gun = secilen_tarih.day
    s_name = str(gun).zfill(2)
    
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = calisma_kitabini_hazirla(yil, ay)
        
    if s_name not in wb.sheetnames:
        ws = wb.create_sheet(title=s_name)
        bos_gunluk_sablon_ciz(ws)
    else:
        ws = wb[s_name]
        
    saat_sayi = int(saat_str.split(":")[0])
    satir = 3 + saat_sayi

    # Önceki Saati Bulup Otomatik Fark Hesaplama
    prev_h = None
    prev_s = None
    if saat_sayi > 0:
        prev_h = ws.cell(row=satir - 1, column=27).value
        prev_s = ws.cell(row=satir - 1, column=29).value
    else:
        # 00:00 saati ise dünün 23:00 saatine bak
        onceki_gun_str = str(gun - 1).zfill(2)
        if onceki_gun_str in wb.sheetnames:
            ws_dun = wb[onceki_gun_str]
            prev_h = ws_dun.cell(row=26, column=27).value
            prev_s = ws_dun.cell(row=26, column=29).value

    havuz_fark = round(float(havuz) - float(prev_h), 2) if (prev_h is not None and str(prev_h).replace('.','',1).replace('-','',1).isdigit()) else 0.00
    sayac_fark = round(float(sayac) - float(prev_s), 3) if (prev_s is not None and str(prev_s).replace('.','',1).isdigit()) else 0.000

    ws.cell(row=satir, column=1, value=sorumlu)
    ws.cell(row=satir, column=2, value=saat_str)

    tum_veriler = u1_vals + u2_vals + u3_vals + [havuz, havuz_fark, sayac, sayac_fark]
    for i, val in enumerate(tum_veriler, start=3):
        ws.cell(row=satir, column=i, value=val)
        
    wb.save(EXCEL_FILE)

st.markdown("---")

# ÜNİTE 1
st.markdown("### 🔵 ÜNİTE 1 (3100 kW)")
c1, c2, c3, c4, c5, c6, c7, c8 = st.columns(8)
with c1: u1_p = st.number_input("Aktif (kW)", value=0.0, step=50.0, key="u1_p")
with c2: u1_q = st.number_input("Reaktif (kVAr)", value=0.0, step=10.0, key="u1_q")
with c3: u1_v = st.number_input("Voltaj (V)", value=6300.0 if u1_p > 0 else 0.0, step=50.0, key="u1_v")
with c4: u1_i = st.number_input("Akım (A)", value=0.0, step=5.0, key="u1_i")
with c5: u1_kanat = st.number_input("Kanat (%)", value=0.0, step=5.0, key="u1_kanat")
with c6: u1_y1 = st.number_input("Ön Yatak 1 (°C)", value=0.0, step=1.0, key="u1_y1")
with c7: u1_y2 = st.number_input("Ön Yatak 2 (°C)", value=0.0, step=1.0, key="u1_y2")
with c8: u1_ay = st.number_input("Arka Yatak (°C)", value=0.0, step=1.0, key="u1_ay")

st.markdown("---")

# ÜNİTE 2
st.markdown("### 🟢 ÜNİTE 2 (3100 kW)")
c1, c2, c3, c4, c5, c6, c7, c8 = st.columns(8)
with c1: u2_p = st.number_input("Aktif (kW)", value=0.0, step=50.0, key="u2_p")
with c2: u2_q = st.number_input("Reaktif (kVAr)", value=0.0, step=10.0, key="u2_q")
with c3: u2_v = st.number_input("Voltaj (V)", value=6300.0 if u2_p > 0 else 0.0, step=50.0, key="u2_v")
with c4: u2_i = st.number_input("Akım (A)", value=0.0, step=5.0, key="u2_i")
with c5: u2_kanat = st.number_input("Kanat (%)", value=0.0, step=5.0, key="u2_kanat")
with c6: u2_y1 = st.number_input("Ön Yatak 1 (°C)", value=0.0, step=1.0, key="u2_y1")
with c7: u2_y2 = st.number_input("Ön Yatak 2 (°C)", value=0.0, step=1.0, key="u2_y2")
with c8: u2_ay = st.number_input("Arka Yatak (°C)", value=0.0, step=1.0, key="u2_ay")

st.markdown("---")

# ÜNİTE 3
st.markdown("### 🔴 ÜNİTE 3 (1200 kW)")
c1, c2, c3, c4, c5, c6, c7, c8 = st.columns(8)
with c1: u3_p = st.number_input("Aktif (kW)", value=0.0, step=50.0, key="u3_p")
with c2: u3_q = st.number_input("Reaktif (kVAr)", value=0.0, step=10.0, key="u3_q")
with c3: u3_v = st.number_input("Voltaj (V)", value=6300.0 if u3_p > 0 else 0.0, step=50.0, key="u3_v")
with c4: u3_i = st.number_input("Akım (A)", value=0.0, step=5.0, key="u3_i")
with c5: u3_kanat = st.number_input("Kanat (%)", value=0.0, step=5.0, key="u3_kanat")
with c6: u3_y1 = st.number_input("Ön Yatak 1 (°C)", value=0.0, step=1.0, key="u3_y1")
with c7: u3_y2 = st.number_input("Ön Yatak 2 (°C)", value=0.0, step=1.0, key="u3_y2")
with c8: u3_ay = st.number_input("Arka Yatak (°C)", value=0.0, step=1.0, key="u3_ay")

st.markdown("---")

# REGÜLATÖR
st.markdown("### 💧 REGÜLATÖR & ANASAYAÇ")
r1, r2 = st.columns(2)
with r1:
    havuz = st.number_input("Yükleme Havuzu Seviyesi (m)", value=958.50, step=0.05, format="%.2f")
    if havuz <= 955.5 and havuz > 0:
        st.error("⚠️ DİKKAT: Havuz seviyesi kritik trip kotuna (955.0 m) çok yakın!")
    elif havuz >= 960.3:
        st.warning("⚠️ DİKKAT: Havuz taşma/savaklama kotuna (960.5 m) yaklaştı!")
with r2:
    sayac = st.number_input("Anasayaç Değeri (2.8.0) (MWh)", value=1500.00, step=0.10, format="%.2f")

st.markdown("---")

# İŞLE BUTONU
kaydet_btn = st.button("💾 VERİYİ SİSTEME VE EXCEL'E İŞLE", use_container_width=True, type="primary")

if kaydet_btn:
    if not sorumlu:
        st.warning("Lütfen vardiya sorumlusu adını girin.")
    else:
        u1_vals = [0.0]*8 if u1_p == 0.0 else [u1_p, u1_q, u1_v, u1_i, u1_kanat, u1_y1, u1_y2, u1_ay]
        u2_vals = [0.0]*8 if u2_p == 0.0 else [u2_p, u2_q, u2_v, u2_i, u2_kanat, u2_y1, u2_y2, u2_ay]
        u3_vals = [0.0]*8 if u3_p == 0.0 else [u3_p, u3_q, u3_v, u3_i, u3_kanat, u3_y1, u3_y2, u3_ay]

        try:
            veriyi_excele_yaz(tarih, secilen_saat, sorumlu, u1_vals, u2_vals, u3_vals, havuz, sayac)
            st.success(f"✅ {secilen_saat} verisi '{EXCEL_FILE}' dosyasındaki '{sayfa_adi}' numaralı güne başarıyla işlendi!")
            
            suanki_idx = saatler_listesi.index(secilen_saat)
            st.session_state["saat_idx"] = (suanki_idx + 1) % 24
            st.rerun()
            
        except PermissionError:
            st.error(f"❌ HATA: '{EXCEL_FILE}' dosyası bilgisayarda açık! Lütfen Excel dosyasını kapatıp butona tekrar basın.")

# --- CANLI MİNİ TEYİT TABLOSU (SEÇİLEN GÜNE AİT) ---
st.markdown("---")
st.markdown(f"### 📋 Canlı Vardiya Kayıt Teyit Tablosu (Dosya: {EXCEL_FILE} | Gün Sayfası: {sayfa_adi})")

if os.path.exists(EXCEL_FILE):
    try:
        wb_oku = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
        if sayfa_adi in wb_oku.sheetnames:
            ws_oku = wb_oku[sayfa_adi]
            headers = [ws_oku.cell(row=2, column=c).value for c in range(1, 31)]
            satirlar = []
            for r in range(3, 27):
                row_vals = [ws_oku.cell(row=r, column=c).value for c in range(1, 31)]
                satirlar.append(row_vals)
            df_goster = pd.DataFrame(satirlar, columns=headers)
            st.dataframe(df_goster, use_container_width=True, height=450)
        else:
            st.info(f"'{sayfa_adi}' sayfası henüz hazır değil.")
    except Exception as e:
        st.warning(f"Tablo okunurken durum oluştu: {e}")
else:
    st.info(f"'{EXCEL_FILE}' çalışma kitabı henüz oluşturulmadı. İlk veriyi kaydettiğinizde ayın tüm günleri hazır olarak açılacaktır.")

# İndirme Butonu
if os.path.exists(EXCEL_FILE):
    with open(EXCEL_FILE, "rb") as f:
        st.download_button(
            label=f"📥 Tüm Aylık Çalışma Kitabını İndir ({EXCEL_FILE})",
            data=f,
            file_name=EXCEL_FILE,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
